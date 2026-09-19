#!/usr/bin/env python3
"""
Edit publications.xlsx by column name - never by position.

Commands
  upgrade             add the README columns that are missing and backfill DOI /
                      Paper Link for existing rows (exact title match on Crossref)
  add FILE.json       append papers (a JSON list of records, the same shape as the
                      "candidates" in .lab-site-updates/new_pubs.json)
  add --doi DOI       append one paper, fetching its metadata from Crossref
  set-section DOI_OR_TITLE SECTION
                      change the Section of an existing row (e.g. a preprint that
                      has now been published)

Every command prints what it changed. Run `python xlsx_to_yml.py --force` afterwards.

Usage
  python scripts/pubs_xlsx.py upgrade
  python scripts/pubs_xlsx.py add .lab-site-updates/approved.json
  python scripts/pubs_xlsx.py add --doi 10.1016/j.clinph.2026.01.001
  python scripts/pubs_xlsx.py set-section 10.31234/osf.io/abcde "Peer-reviewed Journal Paper"
"""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

XLSX = Path(__file__).resolve().parent.parent / "publications.xlsx"

# Column order from the README; existing columns keep their place, missing ones are appended.
README_COLUMNS = [
    "Section", "Authors", "Year", "Date", "Title", "Paper Link", "Journal", "Volume",
    "Issue", "Pages", "DOI", "PDF", "Preprint", "ShareIt", "Supplemental Information",
    "GitHub", "Code", "Data", "Highly Cited", "Hot Paper", "Awards", "Media Coverage",
    "Invited Presentation", "Categories",
]
SECTIONS = {"Peer-reviewed Journal Paper", "Preprint", "Selected Work"}
UA = "ondalab-site-updater/1.0 (lab website maintenance script)"


def norm_title(title: str) -> str:
    t = re.sub(r"<[^>]+>", " ", html.unescape(title or "").lower()).replace("&", " and ")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", t).split())


def venue(item: dict) -> str:
    """Journal name of a Crossref item; for preprints, the server (bioRxiv, PsyArXiv...)."""
    names = (item.get("container-title") or []) + [i.get("name", "") for i in item.get("institution") or []]
    return next((n for n in names if n), "")


def norm_doi(doi: str | None) -> str:
    return re.sub(r"^https?://(dx\.)?doi\.org/", "", (doi or "").strip().lower())


def crossref(path: str) -> dict:
    req = urllib.request.Request("https://api.crossref.org" + path, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=45) as resp:
        return json.loads(resp.read().decode("utf-8"))["message"]


# --------------------------------------------------------------------------

class Sheet:
    def __init__(self, path: Path = XLSX):
        self.path = path
        self.wb = openpyxl.load_workbook(path)
        self.ws = self.wb.active
        self.header = [str(c.value).strip() if c.value is not None else "" for c in self.ws[1]]

    def col(self, name: str) -> int | None:
        """1-based column index for a header name, or None."""
        return self.header.index(name) + 1 if name in self.header else None

    def rows(self):
        for r in range(2, self.ws.max_row + 1):
            if any(self.ws.cell(r, c).value not in (None, "") for c in range(1, len(self.header) + 1)):
                yield r

    def get(self, r: int, name: str):
        c = self.col(name)
        return self.ws.cell(r, c).value if c else None

    def set(self, r: int, name: str, value) -> None:
        c = self.col(name)
        if c is None:
            raise SystemExit(f"Column '{name}' not in publications.xlsx - run `upgrade` first.")
        self.ws.cell(r, c).value = value

    def find(self, key: str) -> int | None:
        """Row whose DOI or normalised title matches key."""
        d, t = norm_doi(key), norm_title(key)
        for r in self.rows():
            if d and norm_doi(str(self.get(r, "DOI") or "")) == d:
                return r
            if t and norm_title(str(self.get(r, "Title") or "")) == t:
                return r
        return None

    def save(self) -> None:
        self.wb.save(self.path)


# --------------------------------------------------------------------------

def cmd_upgrade(_args) -> int:
    sh = Sheet()
    added = [c for c in README_COLUMNS if c not in sh.header]
    for name in added:
        sh.ws.cell(1, len(sh.header) + 1).value = name
        sh.header.append(name)
    print(f"Columns added: {', '.join(added) if added else 'none (already up to date)'}")

    filled, missed = [], []
    for r in sh.rows():
        if sh.get(r, "DOI"):
            continue
        title = str(sh.get(r, "Title") or "")
        q = urllib.parse.urlencode({"query.bibliographic": title, "rows": 5})
        try:
            items = crossref(f"/works?{q}").get("items", [])
        except Exception as exc:  # noqa: BLE001
            missed.append(f"{title[:70]} (Crossref error: {exc})")
            continue
        # Title AND venue must match: the same title often exists as a preprint and as the paper.
        journal = norm_title(str(sh.get(r, "Journal") or ""))
        hit = next((i for i in items
                    if norm_title((i.get("title") or [""])[0]) == norm_title(title)
                    and journal and norm_title(venue(i)) == journal), None)
        if hit:
            sh.set(r, "DOI", hit["DOI"])
            if not sh.get(r, "Paper Link"):
                sh.set(r, "Paper Link", f"https://doi.org/{hit['DOI']}")
            filled.append(f"{title[:70]} -> {hit['DOI']}")
        else:
            near = [f"{i['DOI']} [{venue(i) or i.get('type')}]" for i in items
                    if norm_title((i.get("title") or [""])[0]) == norm_title(title)]
            missed.append(f"{title[:70]} (no title+journal match"
                          + (f"; same title at: {', '.join(near)}" if near else "") + ")")

    sh.save()
    print(f"\nDOIs filled ({len(filled)}):")
    print("\n".join(f"  {x}" for x in filled) or "  -")
    print(f"\nLeft empty, fill by hand if you want them ({len(missed)}):")
    print("\n".join(f"  {x}" for x in missed) or "  -")
    return 0


def record_from_crossref(doi: str) -> dict:
    m = crossref(f"/works/{urllib.parse.quote(doi)}")
    authors = []
    for a in m.get("author", []):
        family = a.get("family", "").strip()
        initials = "".join(p[0] for p in re.split(r"[\s\-.]+", a.get("given", "")) if p)
        if family:
            authors.append(f"{family} {initials}".strip())
    parts = (m.get("published") or m.get("issued") or {}).get("date-parts", [[None]])[0]
    return {
        "title": (m.get("title") or [""])[0],
        "authors": authors,
        "journal": venue(m),
        "year": parts[0],
        "date": "-".join(f"{p:02d}" if i else str(p) for i, p in enumerate(parts) if p),
        "volume": m.get("volume", ""),
        "issue": m.get("issue", ""),
        "pages": m.get("page", "") or m.get("article-number", ""),
        "doi": m.get("DOI", doi),
        "is_preprint": m.get("type") == "posted-content",
    }


def cmd_add(args) -> int:
    if args.doi:
        records = [record_from_crossref(norm_doi(args.doi))]
    elif args.file:
        data = json.loads(Path(args.file).read_text(encoding="utf-8"))
        records = data["candidates"] if isinstance(data, dict) else data
    else:
        raise SystemExit("add needs FILE.json or --doi")

    sh = Sheet()
    missing = [c for c in ("Section", "Authors", "Title", "DOI") if c not in sh.header]
    if missing:
        raise SystemExit(f"publications.xlsx lacks {missing} - run `upgrade` first.")

    added, skipped = [], []
    for rec in records:
        title, doi = rec.get("title", "").strip(), norm_doi(rec.get("doi"))
        existing = (doi and sh.find(doi)) or sh.find(title)
        if existing:
            skipped.append(f"{title[:70]} (already in row {existing}, section: {sh.get(existing, 'Section')})")
            continue
        section = rec.get("section") or ("Preprint" if rec.get("is_preprint") else "Peer-reviewed Journal Paper")
        if section not in SECTIONS:
            raise SystemExit(f"Unknown section '{section}'")
        authors = rec["authors"] if isinstance(rec["authors"], str) else ",".join(a.strip() for a in rec["authors"])
        authors = re.sub(r",\s+", ",", authors)
        year = str(rec.get("year") or "")[:4] or None
        values = {
            "Section": section,
            "Authors": authors,
            "Title": title,
            # PubMed appends subtitles: "Neurological sciences : official journal of ..."
            "Journal": " ".join((rec.get("journal") or "").split(" : ")[0].split()) or None,
            "Year": year,
            "Publication year": year,
            "Date": rec.get("date") or None,
            "Volume": rec.get("volume") or None,
            "Issue": rec.get("issue") or None,
            "Pages": rec.get("pages") or None,
            "DOI": doi or None,
            "Paper Link": f"https://doi.org/{doi}" if doi else None,
        }
        r = sh.ws.max_row + 1
        for name, value in values.items():
            if value is not None and sh.col(name):
                sh.set(r, name, value)
        added.append(f"row {r}: [{section}] {title[:70]}")

    sh.save()
    print(f"Added ({len(added)}):")
    print("\n".join(f"  {x}" for x in added) or "  -")
    if skipped:
        print(f"\nSkipped, already present ({len(skipped)}):")
        print("\n".join(f"  {x}" for x in skipped))
    return 0


def cmd_set_section(args) -> int:
    if args.section not in SECTIONS:
        raise SystemExit(f"Section must be one of {sorted(SECTIONS)}")
    sh = Sheet()
    r = sh.find(args.key)
    if not r:
        raise SystemExit(f"No row matches '{args.key}'")
    old = sh.get(r, "Section")
    sh.set(r, "Section", args.section)
    sh.save()
    print(f"row {r}: Section '{old}' -> '{args.section}'  ({str(sh.get(r, 'Title'))[:70]})")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("upgrade").set_defaults(fn=cmd_upgrade)
    p = sub.add_parser("add")
    p.add_argument("file", nargs="?")
    p.add_argument("--doi")
    p.set_defaults(fn=cmd_add)
    p = sub.add_parser("set-section")
    p.add_argument("key", help="DOI or exact title")
    p.add_argument("section")
    p.set_defaults(fn=cmd_set_section)
    args = ap.parse_args()
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(main())
